import datetime
import json
import re
import sys
from functools import lru_cache
from pathlib import Path

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

FUENTES = Path(__file__).resolve().parent
RAIZ = FUENTES.parent
sys.path.insert(0, str(FUENTES))
import construir_notebooks as cuaderno

CONTENIDO = [FUENTES / 'contenido' / 'teoria.json', FUENTES / 'contenido' / 'taller.json']
FONDOS = FUENTES / 'fondos'
SALIDA = RAIZ / 'presentacion' / 'Minicurso_Analisis_de_Datos.pptx'

TITULO_DOC = 'Introducción al análisis de datos con Python'
ENLACE_TALLER = ''
MARCADOR_ENLACE = '[enlace del taller]'
NOTA_SINTETICOS = 'Datos sintéticos de «Pan de Datos»: base limpia de 1.180 pedidos'
AUTOR = 'Andrés David Galeano Quintero'

VERDE = '06784F'
VERDE_OSCURO = '044E34'
DORADO = 'FFB93E'
TEXTO = '29272D'
SECUNDARIO = '626268'
TARJETA = 'F0F5F2'
BLANCO = 'FFFFFF'
CREMA = 'F7F5F2'
GRIS = '9A9AA0'

X0, X1, Y0, Y1 = 37.4, 682.6, 92.2, 339.8
ANCHO = X1 - X0
LIMITE_PIE = Y1
REDONDEO = 5.0
ESPACIADO = 3.0

ARCHIVOS_FUENTE = {
    ('Calibri', False, False): 'calibri.ttf',
    ('Calibri', True, False): 'calibrib.ttf',
    ('Calibri', False, True): 'calibrii.ttf',
    ('Calibri', True, True): 'calibriz.ttf',
    ('Consolas', False, False): 'consola.ttf',
    ('Consolas', True, False): 'consolab.ttf',
    ('Consolas', False, True): 'consolai.ttf',
    ('Consolas', True, True): 'consolaz.ttf',
}
ALTO_LINEA = {'Calibri': 1.2, 'Consolas': 1.2}
ASCENSO = {'Calibri': 0.2017, 'Consolas': 0.2012}
INSET_REDONDEO = 0.29289
HOLGURA = 0.975
MINIMO_CUERPO = 9.0
VIUDA = 0.16
MARCA = re.compile(r'«(.+?)»')
NBSP = ' '

PROBLEMAS = []
AVISOS = []
ACTUAL = {'id': ''}


def amarrar(texto):
    texto = re.sub(r'«[^»]{1,32}»', lambda m: m.group(0).replace(' ', NBSP), texto)
    texto = re.sub(r"'[^']{1,20}'", lambda m: m.group(0).replace(' ', NBSP), texto)
    texto = re.sub(r'(\d) (?=(min|km|años|pesos|veces)\b)', '\\1' + NBSP, texto)
    texto = re.sub(r'(\d) %', '\\1' + NBSP + '%', texto)
    texto = re.sub(r'(\d) (a|y) (?=\d)', lambda m: m.group(1) + NBSP + m.group(2) + NBSP, texto)
    return texto


def R(texto, tam, color=TEXTO, negrita=False, cursiva=False, fuente='Calibri', spc=0.0):
    if fuente == 'Calibri' and not spc:
        texto = amarrar(texto)
    return {'t': texto, 'tam': tam, 'color': color, 'b': negrita, 'i': cursiva, 'f': fuente, 'spc': spc}


def mixto(texto, tam, tono=TEXTO, negrita=False):
    runs = []
    for k, trozo in enumerate(texto.split('`')):
        if not trozo:
            continue
        if k % 2:
            runs.append(R(trozo, round(tam * 0.92, 2), tono, negrita, fuente='Consolas'))
        else:
            runs.append(R(trozo, tam, tono, negrita))
    return runs


def P(*runs, alin='l', despues=0.0, inter=1.0, vineta=False):
    return {'runs': list(runs), 'alin': alin, 'despues': despues, 'inter': inter, 'vineta': vineta}


@lru_cache(None)
def fuente_pil(nombre, negrita, cursiva):
    return ImageFont.truetype(str(Path('C:/Windows/Fonts') / ARCHIVOS_FUENTE[(nombre, negrita, cursiva)]), 1000)


@lru_cache(None)
def ancho_caracter(c, nombre, negrita, cursiva):
    return fuente_pil(nombre, negrita, cursiva).getlength(c) / 1000


def anchos(parrafo, escala):
    res = []
    for r in parrafo['runs']:
        for c in r['t']:
            res.append(ancho_caracter(c, r['f'], r['b'], r['i']) * r['tam'] * escala + r['spc'])
    return res


def ancho_linea(parrafo, escala=1.0):
    return sum(anchos(parrafo, escala))


def contar_lineas(parrafo, ancho, escala=1.0):
    texto = ''.join(r['t'] for r in parrafo['runs'])
    if not texto.strip():
        return 1
    if parrafo['vineta']:
        ancho -= 14
    ancho *= HOLGURA
    w = anchos(parrafo, escala)
    n, actual, pendiente = 1, 0.0, 0.0
    for m in re.finditer(r'(\s*\S+)(\s*)', texto):
        palabra = sum(w[m.start(1):m.end(1)])
        espacio = sum(w[m.start(2):m.end(2)])
        if actual == 0.0:
            actual = palabra
        elif actual + pendiente + palabra <= ancho:
            actual += pendiente + palabra
        else:
            n += 1
            actual = palabra
        while actual > ancho:
            n += 1
            actual -= ancho
        pendiente = espacio
    return n


def alto_parrafo(parrafo, ancho, escala=1.0):
    tam = max(r['tam'] for r in parrafo['runs']) * escala
    factor = max(ALTO_LINEA[r['f']] for r in parrafo['runs'])
    return contar_lineas(parrafo, ancho, escala) * tam * factor * parrafo['inter'] + parrafo['despues']


def alto_texto(parrafos, ancho, escala=1.0):
    total = sum(alto_parrafo(p, ancho, escala) for p in parrafos)
    return total - parrafos[-1]['despues']


def ajustar(parrafos, ancho, alto, donde, minimo=MINIMO_CUERPO):
    menor = min(r['tam'] for p in parrafos for r in p['runs'])
    escala = 1.0
    while escala * menor >= minimo - 1e-6:
        if alto_texto(parrafos, ancho, escala) <= alto:
            return escala
        escala = round(escala - 0.01, 2)
    PROBLEMAS.append(f'{donde}: el texto no cabe en {ancho:.0f} × {alto:.0f} pt ni con {minimo} pt '
                     f'(necesita {alto_texto(parrafos, ancho, minimo / menor):.0f} pt)')
    return minimo / menor


def color(hexa):
    return RGBColor.from_string(hexa)


def escribir(marco, parrafos, escala=1.0):
    marco.word_wrap = True
    marco.auto_size = MSO_AUTO_SIZE.NONE
    for k, p in enumerate(parrafos):
        par = marco.paragraphs[0] if k == 0 else marco.add_paragraph()
        par.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[p['alin']]
        if p['inter'] != 1.0:
            par.line_spacing = round(p['inter'], 3)
        if p['despues'] and k < len(parrafos) - 1:
            par.space_after = Pt(round(p['despues'], 1))
        if p['vineta']:
            ppr = par._p.get_or_add_pPr()
            ppr.set('marL', str(Pt(14)))
            ppr.set('indent', str(-Pt(14)))
            bu = OxmlElement('a:buChar')
            bu.set('char', '•')
            ppr.append(bu)
        for r in p['runs']:
            run = par.add_run()
            run.text = r['t']
            f = run.font
            f.name = r['f']
            f.size = Pt(round(r['tam'] * escala, 1))
            f.bold = r['b']
            f.italic = r['i']
            f.color.rgb = color(r['color'])
            rpr = run._r.get_or_add_rPr()
            rpr.set('lang', 'es-CO')
            if r['spc']:
                rpr.set('spc', str(int(round(r['spc'] * 100))))


def preparar_marco(forma, margen=(0, 0, 0, 0), ancla='t'):
    marco = forma.text_frame
    marco.margin_left, marco.margin_top, marco.margin_right, marco.margin_bottom = [Pt(m) for m in margen]
    marco.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[ancla]
    return marco


def caja(diapo, x, y, w, h, parrafos, escala=1.0, ancla='t', nombre=None):
    forma = diapo.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    escribir(preparar_marco(forma, ancla=ancla), parrafos, escala)
    if nombre:
        forma.name = nombre
    return forma


def rectangulo(diapo, x, y, w, h, relleno, redondeo=0.0, nombre=None):
    tipo = MSO_SHAPE.ROUNDED_RECTANGLE if redondeo else MSO_SHAPE.RECTANGLE
    forma = diapo.shapes.add_shape(tipo, Pt(x), Pt(y), Pt(w), Pt(h))
    if redondeo:
        forma.adjustments[0] = min(0.5, redondeo / min(w, h))
    forma.fill.solid()
    forma.fill.fore_color.rgb = color(relleno)
    forma.line.fill.background()
    forma.shadow.inherit = False
    if nombre:
        forma.name = nombre
    return forma


def bloque(diapo, x, y, w, h, relleno, parrafos, escala=1.0, margen=(12, 12, 12, 12), ancla='t',
           redondeo=REDONDEO, nombre=None):
    forma = rectangulo(diapo, x, y, w, h, relleno, redondeo, nombre)
    escribir(preparar_marco(forma, margen, ancla), parrafos, escala)
    return forma


def flecha(diapo, x1, y1, x2, y2, tono=VERDE, grosor=1.25):
    linea = diapo.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    linea.line.color.rgb = color(tono)
    linea.line.width = Pt(grosor)
    punta = OxmlElement('a:tailEnd')
    punta.set('type', 'triangle')
    punta.set('w', 'med')
    punta.set('len', 'med')
    linea.line._get_or_add_ln().append(punta)
    return linea


def figura(diapo, ruta, x, y, w_max, h_max, alin='c', valin='t', descripcion=''):
    archivo = FUENTES / ruta
    if not archivo.exists():
        PROBLEMAS.append(f'no existe la figura {ruta}')
        return x, y, 0, 0
    with Image.open(archivo) as im:
        proporcion = im.width / im.height
    w = min(w_max, h_max * proporcion)
    h = w / proporcion
    fx = {'l': x, 'c': x + (w_max - w) / 2, 'r': x + w_max - w}[alin]
    fy = {'t': y, 'm': y + (h_max - h) / 2, 'b': y + h_max - h}[valin]
    imagen = diapo.shapes.add_picture(str(archivo), Pt(fx), Pt(fy), Pt(w), Pt(h))
    imagen.name = Path(ruta).stem
    if descripcion:
        imagen._element.nvPicPr.cNvPr.set('descr', descripcion)
    return fx, fy, w, h


def y_texto(y_letra, tam, fuente='Calibri'):
    return y_letra - ASCENSO[fuente] * tam


def nueva(prs, fondo):
    diapo = prs.slides.add_slide(prs.slide_layouts[6])
    fondo_img = diapo.shapes.add_picture(str(FONDOS / fondo), 0, 0, prs.slide_width, prs.slide_height)
    fondo_img.name = 'Fondo'
    return diapo


def notas(diapo, d):
    marco = diapo.notes_slide.notes_text_frame
    marco.text = d['notas']
    for par in marco.paragraphs:
        for run in par.runs:
            run._r.get_or_add_rPr().set('lang', 'es-CO')


def numero_pagina(diapo, numero):
    caja(diapo, 21.6, y_texto(6.7, 10), 40, 14, [P(R(str(numero), 10, SECUNDARIO))], nombre='Número')


def etiqueta_runs(texto, tam, tono):
    return P(R(texto, tam, tono, spc=ESPACIADO))


def titulo_contenido(diapo, d):
    tam = 21.0
    par = P(R(d['titulo'], tam, VERDE))
    while ancho_linea(par, tam / 21.0) > ANCHO * HOLGURA and tam > 18:
        tam -= 0.5
    if ancho_linea(par, tam / 21.0) > ANCHO * HOLGURA:
        PROBLEMAS.append(f"{d['id']}: el título no cabe en una línea")
    caja(diapo, X0 + 0.1, y_texto(37.3, 21), ANCHO, 27, [P(R(d['titulo'], tam, VERDE))], nombre='Título')
    etq = etiqueta_runs(d['etiqueta'], 10, VERDE)
    if ancho_linea(etq) > ANCHO:
        PROBLEMAS.append(f"{d['id']}: la etiqueta no cabe en una línea")
    caja(diapo, X0 + 0.1, y_texto(73.0, 10), ANCHO, 14, [etq], nombre='Etiqueta')


def base_contenido(prs, d, numero):
    diapo = nueva(prs, 'fondo_contenido.png')
    numero_pagina(diapo, numero)
    titulo_contenido(diapo, d)
    caja(diapo, 20.2, y_texto(373.7, 9), 300, 12,
         [P(R('Facultad de Minas  ·  Sede Medellín', 9, VERDE, cursiva=True))], nombre='Pie')
    return diapo


def pie_verde(diapo):
    caja(diapo, X0 + 0.1, y_texto(373.7, 10.6), 500, 14,
         [P(R('Universidad Nacional de Colombia  ·  Facultad de Minas  ·  Sede Medellín', 10.6, BLANCO))],
         nombre='Pie')


def portada(prs, d, numero):
    diapo = nueva(prs, 'fondo_portada.png')
    lineas = d['titulo'].split('\n')
    inter = 36.4 / (27 * ALTO_LINEA['Calibri'])
    pars = [P(R(t, 27, VERDE), inter=inter) for t in lineas]
    for p in pars:
        if ancho_linea(p) > ANCHO * HOLGURA:
            PROBLEMAS.append(f"{d['id']}: una línea del título no cabe")
    alto = alto_texto(pars, ANCHO)
    caja(diapo, X0 + 0.1, y_texto(118.6, 27) - 27 * ALTO_LINEA['Calibri'] * (inter - 1), ANCHO, alto + 4, pars,
         nombre='Título')
    y = 118.6 + 36.4 * (len(lineas) - 1) + 27 + 37.4
    caja(diapo, X0 + 0.1, y_texto(y, 10), ANCHO, 14, [etiqueta_runs(d['etiqueta'], 10, VERDE)], nombre='Etiqueta')
    caja(diapo, X0 + 0.1, y_texto(y + 43.8, 13), ANCHO, 18, [P(R(d['autores'], 13, TEXTO))], nombre='Autores')
    nota = [P(R(d['nota'], 10.6, SECUNDARIO))]
    caja(diapo, X0 + 0.1, y_texto(y + 63.8, 10.6), ANCHO, alto_texto(nota, ANCHO) + 4, nota, nombre='Nota')
    pie_verde(diapo)
    return diapo


def seccion(prs, d, numero):
    diapo = nueva(prs, 'fondo_portada.png')
    numero_pagina(diapo, numero)
    tit = P(R(d['titulo'], 27, VERDE))
    if ancho_linea(tit) > ANCHO * HOLGURA:
        PROBLEMAS.append(f"{d['id']}: el título no cabe en una línea")
    caja(diapo, X0 + 0.1, y_texto(136.0, 27), ANCHO, 36, [tit], nombre='Título')
    caja(diapo, X0 + 0.1, y_texto(136.0 + 27 + 37.4, 10), ANCHO, 14,
         [etiqueta_runs(d['etiqueta'], 10, VERDE)], nombre='Etiqueta')
    texto = [P(R(d['texto'], 13, TEXTO), inter=1.1)]
    w = 470
    escala = ajustar(texto, w, 60, d['id'])
    caja(diapo, X0 + 0.1, y_texto(136.0 + 27 + 37.4 + 36.6, 13), w, 62, texto, escala, nombre='Texto')
    pie_verde(diapo)
    return diapo


def parrafos_tarjeta(titulo, texto, tam_titulo=12.0, tam_texto=9.8):
    pars = []
    if titulo:
        pars.append(P(R(titulo, tam_titulo, VERDE), despues=5.5))
    pars.append(P(R(texto, tam_texto, TEXTO), inter=13.2 / (9.8 * ALTO_LINEA['Calibri'])))
    return pars


def parrafos_destacado(etiqueta, texto, tam=12.6):
    pars = []
    if etiqueta:
        pars.append(P(R(etiqueta, 9.6, DORADO, spc=ESPACIADO), despues=11.2))
    pars.append(P(*mixto(texto, tam, BLANCO), inter=17.5 / (12.6 * ALTO_LINEA['Calibri'])))
    return pars


MARGEN_TARJETA = (13.7, 12.0, 12.0, 12.0)
MARGEN_DESTACADO = (20.2, 12.7, 20.2, 12.0)


def alto_bloque(pars, w, margen, escala=1.0):
    return alto_texto(pars, w - margen[0] - margen[2], escala) + margen[1] + margen[3]


def fila_tarjetas(diapo, d, contenidos, y, alto_min, alto_max, nombre):
    n = len(contenidos)
    hueco = 15.8 if n < 4 else 13.0
    w = (ANCHO - hueco * (n - 1)) / n
    interior = w - MARGEN_TARJETA[0] - MARGEN_TARJETA[2]
    escala = 1.0
    while True:
        alto = max(alto_texto(c, interior, escala) for c in contenidos) + MARGEN_TARJETA[1] + MARGEN_TARJETA[3]
        menor = min(r['tam'] for c in contenidos for p in c for r in p['runs'])
        if alto <= alto_max or escala * menor <= MINIMO_CUERPO:
            break
        escala = round(escala - 0.01, 2)
    if alto > alto_max:
        PROBLEMAS.append(f"{d['id']}: las tarjetas necesitan {alto:.0f} pt y hay {alto_max:.0f}")
    alto = max(alto, alto_min)
    for k, c in enumerate(contenidos):
        bloque(diapo, X0 + k * (w + hueco), y, w, alto, TARJETA, c, escala, MARGEN_TARJETA,
               nombre=f'{nombre} {k + 1}')
    return alto


def tarjetas(prs, d, numero):
    diapo = base_contenido(prs, d, numero)
    contenidos = [parrafos_tarjeta(t['titulo'], t['texto']) for t in d['tarjetas']]
    hueco = 18.8
    dest = d.get('destacado')
    alto_dest = 0.0
    pars_dest = None
    if dest:
        pars_dest = parrafos_destacado(dest['etiqueta'], dest['texto'])
        alto_dest = alto_bloque(pars_dest, ANCHO, MARGEN_DESTACADO)
    disponible = Y1 - Y0 - (alto_dest + hueco if dest else 0)
    n = len(contenidos)
    w = (ANCHO - (15.8 if n < 4 else 13.0) * (n - 1)) / n
    necesario = max(alto_texto(c, w - 25.7) for c in contenidos) + 24
    sobra = max(0.0, disponible - necesario)
    extra = min(sobra / 2, 28.0) if dest else min(sobra, 20.0)
    alto = fila_tarjetas(diapo, d, contenidos, Y0, necesario + extra, disponible, 'Tarjeta')
    if dest:
        y = Y0 + alto + hueco
        alto_d = max(alto_dest, min(Y1 - y, alto_dest + 28.0))
        bloque(diapo, X0, y, ANCHO, alto_d, VERDE, pars_dest, 1.0, MARGEN_DESTACADO, ancla='m', nombre='Destacado')
    return diapo


def cifras(prs, d, numero):
    diapo = base_contenido(prs, d, numero)
    contenidos = []
    n = len(d['cifras'])
    hueco_t = 15.8 if n < 4 else 13.0
    w = (ANCHO - hueco_t * (n - 1)) / n
    for c in d['cifras']:
        tam = 26.0
        valor = P(R(c['valor'], tam, VERDE), despues=4)
        while ancho_linea(valor) > (w - 25.7) * HOLGURA and tam > 20:
            tam -= 1
            valor = P(R(c['valor'], tam, VERDE), despues=4)
        contenidos.append([valor, P(R(c['leyenda'], 9.6, SECUNDARIO), inter=1.0)])
    hueco = 18.8
    pars_dest = None
    alto_dest = 0.0
    if d.get('texto'):
        pars_dest = parrafos_destacado(None, d['texto'])
        alto_dest = alto_bloque(pars_dest, ANCHO, MARGEN_DESTACADO)
    disponible = Y1 - Y0 - (alto_dest + hueco if pars_dest else 0)
    necesario = max(alto_texto(c, w - 25.7) for c in contenidos) + 24
    extra = min(max(0.0, disponible - necesario) / 2, 22.0)
    alto = fila_tarjetas(diapo, d, contenidos, Y0, necesario + extra, disponible, 'Cifra')
    y = Y0 + alto + hueco
    if pars_dest:
        alto_d = max(alto_dest, min(Y1 - y, alto_dest + 22.0))
        bloque(diapo, X0, y, ANCHO, alto_d, VERDE, pars_dest, 1.0, MARGEN_DESTACADO, ancla='m', nombre='Destacado')
    elif d.get('tabla'):
        dibujar_tabla(diapo, d, d['tabla']['columnas'], d['tabla']['filas'], X0, y, ANCHO, Y1 - y)
    return diapo


def asignar_anchos(naturales, total, minimos):
    n = len(naturales)
    if sum(naturales) <= total:
        extra = total - sum(naturales)
        return [a + extra * a / sum(naturales) for a in naturales]
    fijos = {}
    while True:
        resto = total - sum(fijos.values())
        libres = [k for k in range(n) if k not in fijos]
        suma = sum(naturales[k] for k in libres)
        cambio = False
        for k in libres:
            parte = resto * naturales[k] / suma
            if naturales[k] <= parte:
                fijos[k] = naturales[k]
                cambio = True
            elif parte < minimos[k]:
                fijos[k] = minimos[k]
                cambio = True
        if not cambio:
            return [fijos[k] if k in fijos else resto * naturales[k] / suma for k in range(n)]


def celda_par(texto, tam, tono, negrita=False, fuente='Calibri'):
    return P(R(texto, tam, tono, negrita, fuente=fuente))


def dibujar_tabla(diapo, d, columnas, filas, x, y, w, alto_max, columnas_codigo=()):
    ncol = len(columnas)
    margen_x = 6.5
    n_filas = len(filas) + 1
    tam_cab = 9.6
    tam = 10.6 if n_filas <= 7 else 10.4
    while True:
        cab = [celda_par(c, tam_cab, BLANCO, True) for c in columnas]
        cuerpo = [[celda_par(v, tam, TEXTO, fuente='Consolas' if j in columnas_codigo else 'Calibri')
                   for j, v in enumerate(fila)] for fila in filas]
        naturales = []
        minimos = []
        for j in range(ncol):
            todos = [cab[j]] + [f[j] for f in cuerpo]
            naturales.append(max(ancho_linea(p) for p in todos) / HOLGURA + 2 * margen_x + 1)
            palabra = max(ancho_linea(P(dict(r, t=pal))) for p in todos for r in p['runs'] for pal in r['t'].split())
            minimos.append(palabra / HOLGURA + 2 * margen_x + 1)
        anchos_col = asignar_anchos(naturales, w, minimos)
        margen_y = 4.0 if n_filas <= 7 else (3.2 if n_filas <= 10 else 2.8)
        altos = []
        for fila in [cab] + cuerpo:
            altos.append(max(alto_texto([p], anchos_col[j] - 2 * margen_x) for j, p in enumerate(fila)) + 2 * margen_y)
        objetivo = min(34.0, alto_max / n_filas) if n_filas <= 7 else 0
        altos = [max(a, objetivo) for a in altos]
        total = sum(altos)
        if total <= alto_max or tam <= MINIMO_CUERPO:
            break
        tam = round(tam - 0.2, 1)
    if total > alto_max:
        PROBLEMAS.append(f"{d['id']}: la tabla necesita {total:.0f} pt y hay {alto_max:.0f}")
    forma = diapo.shapes.add_table(n_filas, ncol, Pt(x), Pt(y), Pt(w), Pt(total))
    forma.name = 'Tabla'
    tabla = forma.table
    tbl_pr = forma._element.graphic.graphicData.tbl.tblPr
    tbl_pr.set('firstRow', '0')
    tbl_pr.set('bandRow', '0')
    estilo = tbl_pr.find(qn('a:tableStyleId'))
    if estilo is None:
        estilo = OxmlElement('a:tableStyleId')
        tbl_pr.append(estilo)
    estilo.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    for j, a in enumerate(anchos_col):
        tabla.columns[j].width = Pt(a)
    for i, a in enumerate(altos):
        tabla.rows[i].height = Pt(a)
    for i, fila in enumerate([cab] + cuerpo):
        relleno = VERDE if i == 0 else (BLANCO if i % 2 == 1 else TARJETA)
        for j, p in enumerate(fila):
            celda = tabla.cell(i, j)
            celda.fill.solid()
            celda.fill.fore_color.rgb = color(relleno)
            celda.margin_left = Pt(margen_x)
            celda.margin_right = Pt(margen_x)
            celda.margin_top = Pt(margen_y)
            celda.margin_bottom = Pt(margen_y)
            celda.vertical_anchor = MSO_ANCHOR.MIDDLE
            marco = celda.text_frame
            marco.word_wrap = True
            escribir(marco, [p])
    return total


def parrafos_nota(texto, tam):
    return [P(*mixto(texto, tam, TEXTO), inter=1.05)]


def nota_inferior(diapo, d, texto, y, w=ANCHO, x=X0, limite=LIMITE_PIE, tam=10.6):
    pars = parrafos_nota(texto, tam)
    alto_disp = limite - y
    escala = ajustar(pars, w, alto_disp, f"{d['id']} (nota)")
    alto = alto_texto(pars, w, escala)
    caja(diapo, x + 0.1, y, w, alto + 3, pars, escala, nombre='Nota')
    return alto


def tabla(prs, d, numero):
    diapo = base_contenido(prs, d, numero)
    nota = d.get('nota')
    alto_nota = 0.0
    w_tabla = ANCHO
    if d.get('figura'):
        w_tabla = ANCHO * 0.6
    if nota:
        alto_nota = alto_texto(parrafos_nota(nota, 10.6), w_tabla) + 10
    disponible = Y1 - Y0 - max(0.0, alto_nota - (LIMITE_PIE - Y1))
    total = dibujar_tabla(diapo, d, d['columnas'], d['filas'], X0, Y0, w_tabla, disponible,
                          d.get('columnas_codigo', ()))
    if nota:
        nota_inferior(diapo, d, nota, Y0 + total + 10, w_tabla)
    if d.get('figura'):
        x = X0 + w_tabla + 18
        pie = d.get('pie')
        alto_pie = alto_texto([P(R(pie, 8.6, SECUNDARIO, cursiva=True))], X1 - x) + 6 if pie else 0
        fx, fy, fw, fh = figura(diapo, d['figura'], x, Y0, X1 - x, Y1 - Y0 - alto_pie, 'c', 'm', pie or '')
        if pie:
            caja(diapo, fx, fy + fh + 6, X1 - fx, alto_pie, [P(R(pie, 8.6, SECUNDARIO, cursiva=True))], nombre='Pie de figura')
    return diapo


def parrafos_texto(d, tam=11.0):
    pars = []
    if d.get('parrafos'):
        for p in d['parrafos']:
            runs = []
            if p.get('negrita'):
                runs.append(R(p['negrita'] + ' ', tam, VERDE_OSCURO, True))
            runs.append(R(p['texto'], tam, TEXTO))
            pars.append(P(*runs, despues=9.0, inter=1.08))
    for v in d.get('vinetas', []):
        pars.append(P(R(v, tam, TEXTO), despues=7.0, inter=1.08, vineta=True))
    return pars


def texto_figura(prs, d, numero):
    diapo = base_contenido(prs, d, numero)
    pie = d.get('pie')
    pie_par = [P(R(pie, 8.6, SECUNDARIO, cursiva=True), inter=1.05)] if pie else None
    if d.get('figura_ancha'):
        n = len(d['parrafos'])
        hueco = 22.0
        w_col = (ANCHO - hueco * (n - 1)) / n
        columnas = [parrafos_texto({'parrafos': [p]}) for p in d['parrafos']]
        escala = min(ajustar(c, w_col, 60, d['id']) for c in columnas)
        alto_txt = max(alto_texto(c, w_col, escala) for c in columnas)
        for k, c in enumerate(columnas):
            caja(diapo, X0 + 0.1 + k * (w_col + hueco), Y0 + 2, w_col, alto_txt + 4, c, escala, nombre=f'Texto {k + 1}')
        y = Y0 + 2 + alto_txt + 14
        alto_pie = alto_texto(pie_par, ANCHO) + 6 if pie else 0
        fx, fy, fw, fh = figura(diapo, d['figura'], X0, y, ANCHO, Y1 - y - alto_pie, 'c', 't', pie or '')
        if pie:
            caja(diapo, fx, fy + fh + 6, fw, alto_pie, pie_par, nombre='Pie de figura')
        return diapo
    w_texto = 300.0
    x_fig = X0 + w_texto + 18
    w_fig = X1 - x_fig
    alto_pie = alto_texto(pie_par, w_fig) + 6 if pie else 0
    h_fig_max = Y1 - Y0 - alto_pie
    with Image.open(FUENTES / d['figura']) as im:
        h_fig = min(h_fig_max, w_fig * im.height / im.width)
    y_fig = Y0 + 2
    fx, fy, fw, fh = figura(diapo, d['figura'], x_fig, y_fig, w_fig, h_fig, 'r', 't', pie or '')
    if pie:
        caja(diapo, fx, fy + fh + 6, fw, alto_pie, pie_par, nombre='Pie de figura')
    pars = parrafos_texto(d)
    escala = ajustar(pars, w_texto, Y1 - Y0, d['id'])
    caja(diapo, X0 + 0.1, Y0 + 2, w_texto, Y1 - Y0, pars, escala, nombre='Texto')
    return diapo


def flujo(prs, d, numero):
    diapo = base_contenido(prs, d, numero)
    pasos = d['pasos']
    n = len(pasos)
    hueco = 18.0
    w = (ANCHO - hueco * (n - 1)) / n
    resaltar = set(d.get('resaltar', []))
    tam_txt = 11.0
    margen = (9.0, 11.0, 8.0, 11.0)
    interior = w - margen[0] - margen[2]
    tam_cod = 10.5
    todos = [c for p in pasos for c in p.get('codigo', [])]
    while todos and tam_cod > MINIMO_CUERPO and max(
            ancho_linea(P(R(c, tam_cod, negrita=True, fuente='Consolas'))) for c in todos) > interior * HOLGURA:
        tam_cod = round(tam_cod - 0.25, 2)
    contenidos = []
    codigos = []
    for k, p in enumerate(pasos, start=1):
        claro = k in resaltar
        contenidos.append([P(R(p['titulo'], 13, BLANCO if claro else VERDE), despues=6.0),
                           P(R(p['texto'], tam_txt, BLANCO if claro else TEXTO),
                             inter=14.7 / (tam_txt * ALTO_LINEA['Calibri']))])
        codigos.append([P(R(c, tam_cod, DORADO if claro else VERDE, True, fuente='Consolas'), despues=2.0)
                        for c in p.get('codigo', [])])
    alto_codigo = max((alto_texto(c, interior) + 8 for c in codigos if c), default=0.0)
    reserva = (21.0 if d.get('ciclo') else 0.0)
    if d.get('nota'):
        reserva += 20 + alto_texto(parrafos_nota(d['nota'], 11.5), ANCHO)
    alto_max = min(175.0, Y1 - (Y0 + 8) - reserva - 4)
    escala = 1.0
    while True:
        alto = max(alto_texto(c, interior, escala) for c in contenidos) + margen[1] + margen[3] + alto_codigo
        if alto <= alto_max or escala * tam_txt <= MINIMO_CUERPO:
            break
        escala = round(escala - 0.01, 2)
    if alto > alto_max:
        PROBLEMAS.append(f"{d['id']}: las cajas del flujo necesitan {alto:.0f} pt y hay {alto_max:.0f}")
    alto = max(alto, min(alto_max, alto + 30.0))
    y = Y0 + 8
    for k, c in enumerate(contenidos, start=1):
        x = X0 + (k - 1) * (w + hueco)
        relleno = VERDE if k in resaltar else TARJETA
        bloque(diapo, x, y, w, alto, relleno, c, escala, margen, nombre=f'Paso {k}')
        if codigos[k - 1]:
            alto_c = alto_texto(codigos[k - 1], interior)
            caja(diapo, x + margen[0], y + alto - margen[3] - alto_c - 2, interior, alto_c + 2, codigos[k - 1],
                 nombre=f'Método {k}')
        if k < n:
            flecha(diapo, x + w + 3, y + alto / 2, x + w + hueco - 3, y + alto / 2)
    abajo = y + alto
    if d.get('ciclo'):
        x_ini = X0 + w / 2
        x_fin = X0 + (n - 1) * (w + hueco) + w / 2
        y_ret = abajo + 18
        forma = diapo.shapes.build_freeform(Pt(x_fin), Pt(abajo + 3))
        forma.add_line_segments([(Pt(x_fin), Pt(y_ret)), (Pt(x_ini), Pt(y_ret)), (Pt(x_ini), Pt(abajo + 3))],
                                close=False)
        retorno = forma.convert_to_shape()
        retorno.name = 'Regreso al inicio'
        retorno.fill.background()
        retorno.line.color.rgb = color(GRIS)
        retorno.line.width = Pt(1.25)
        retorno.shadow.inherit = False
        punta = OxmlElement('a:tailEnd')
        punta.set('type', 'triangle')
        punta.set('w', 'med')
        punta.set('len', 'med')
        retorno.line._get_or_add_ln().append(punta)
        abajo = y_ret
    if d.get('nota'):
        nota_inferior(diapo, d, d['nota'], abajo + 20, tam=11.5)
    return diapo


def codigo_marcado(d):
    celdas = [c[1] for c in cuaderno.celdas_de_codigo()]
    for celda in celdas:
        if cuaderno.solucion(celda) == d['codigo']:
            return celda
    buscadas = d['codigo'].split('\n')
    marcadas = []
    i = 0
    while i < len(buscadas):
        mejor = []
        for celda in celdas:
            lineas = celda.split('\n')
            for inicio in range(len(lineas)):
                k = 0
                while (inicio + k < len(lineas) and i + k < len(buscadas)
                       and cuaderno.solucion(lineas[inicio + k]) == buscadas[i + k]):
                    k += 1
                if k > len(mejor):
                    mejor = lineas[inicio:inicio + k]
        if not mejor:
            PROBLEMAS.append(f"{d['id']}: el código no coincide con las celdas de construir_notebooks.py")
            return d['codigo']
        marcadas += mejor
        i += len(mejor)
    return '\n'.join(marcadas)


def resaltar_parametros(linea):
    return re.sub(r"(\w+=')«([^»]{1,2})»(')", r'«\1\2\3»', linea)


def parrafos_codigo(d, tam):
    pars = []
    for linea in codigo_marcado(d).split('\n'):
        linea = resaltar_parametros(linea)
        runs = []
        for k, trozo in enumerate(MARCA.split(linea)):
            if not trozo:
                continue
            if k % 2:
                runs.append(R(trozo, tam, VERDE, True, fuente='Consolas'))
            else:
                runs.append(R(trozo, tam, TEXTO, fuente='Consolas'))
        pars.append(P(*runs))
    return pars


MARGEN_CODIGO = (14.0, 8.0, 10.0, 8.0)


def caja_codigo(diapo, d, x, y, w, tam_max=12.0):
    tam = tam_max
    while True:
        pars = parrafos_codigo(d, tam)
        largo = max(ancho_linea(p) for p in pars)
        if largo <= (w - MARGEN_CODIGO[0] - MARGEN_CODIGO[2]) * 0.985 or tam <= 10.5:
            break
        tam = round(tam - 0.25, 2)
    if largo > (w - MARGEN_CODIGO[0] - MARGEN_CODIGO[2]):
        PROBLEMAS.append(f"{d['id']}: una línea de código no cabe a {tam} pt")
    alto = len(pars) * tam * ALTO_LINEA['Consolas'] + MARGEN_CODIGO[1] + MARGEN_CODIGO[3]
    bloque(diapo, x, y, w, alto, TARJETA, pars, 1.0, MARGEN_CODIGO, redondeo=0, nombre='Código')
    rectangulo(diapo, x, y, 3, alto, VERDE, nombre='Barra del código')
    return alto, tam


def columnas_ejemplos(valores, tam, interior):
    if len(valores) <= 3:
        return 1
    izquierda = max(ancho_linea(P(R(v, tam, fuente='Consolas'))) for v in valores[0::2])
    derecha = max(ancho_linea(P(R(v, tam, fuente='Consolas'))) for v in valores[1::2])
    return 2 if izquierda <= interior / 2 - 16 and derecha <= interior / 2 else 1


def tarjeta_ejemplos(diapo, d, x, y, w, alto, calcular=False):
    margen = (12.0, 10.0, 10.0, 10.0)
    tam = 10.5
    valores = d['ejemplo']
    interior = w - margen[0] - margen[2]
    mayor = max(ancho_linea(P(R(v, tam, fuente='Consolas'))) for v in valores)
    columnas = columnas_ejemplos(valores, tam, interior)
    if mayor > interior * HOLGURA:
        PROBLEMAS.append(f"{d['id']}: un ejemplo no cabe en la tarjeta")
    filas = -(-len(valores) // columnas)
    etiqueta = [P(R('EJEMPLOS', 9.0, VERDE, spc=ESPACIADO))]
    alto_etq = alto_texto(etiqueta, interior) + 6
    paso = tam * ALTO_LINEA['Consolas'] + 2.5
    necesario = margen[1] + alto_etq + filas * paso - 2.5 + margen[3]
    if calcular:
        return necesario
    rectangulo(diapo, x, y, w, alto, TARJETA, REDONDEO, nombre='Ejemplos')
    caja(diapo, x + margen[0], y + margen[1], interior, alto_etq, etiqueta, nombre='Etiqueta de ejemplos')
    ancho_col = interior / columnas
    for c in range(columnas):
        pars = [P(R(v, tam, TEXTO, fuente='Consolas'), despues=2.5) for v in valores[c::columnas]]
        caja(diapo, x + margen[0] + c * ancho_col, y + margen[1] + alto_etq, ancho_col, filas * paso + 2, pars,
             nombre=f'Valores {c + 1}')
    return necesario


W_DER_ETL = 240.0
HUECO_ETL = 18.0
SEPARACION_ETL = 14.0
TAM_CODIGO_ETL = 12.0
ETL = {}


def parrafos_paso(d):
    pars = [P(R('Problema: ', 11.5, VERDE_OSCURO, True), *mixto(d['problema'], 11.5), despues=8, inter=1.05),
            P(R('Diagnóstico: ', 11.5, VERDE_OSCURO, True), R(d['diagnostico'], 10.5, TEXTO, fuente='Consolas'),
              despues=3)]
    if d.get('comprobacion'):
        pars.append(P(R('Comprobación: ', 11.5, VERDE_OSCURO, True),
                      R(d['comprobacion'], 10.5, TEXTO, fuente='Consolas'), despues=3))
    pars.append(P(R('Método clave: ', 11.5, VERDE_OSCURO, True), R(d['metodo'], 10.5, VERDE, True, fuente='Consolas')))
    return pars


def alto_superior_paso(d):
    w_izq = ANCHO - W_DER_ETL - HUECO_ETL
    return max(alto_texto(parrafos_paso(d), w_izq), tarjeta_ejemplos(None, d, 0, 0, W_DER_ETL, 0, calcular=True))


def preparar_pasos(diapositivas):
    pasos = [d for d in diapositivas if d['layout'] == 'paso_etl']
    if pasos:
        ETL['alto_sup'] = max(alto_superior_paso(d) for d in pasos)
        ETL['y_codigo'] = Y0 + ETL['alto_sup'] + SEPARACION_ETL


def paso_etl(prs, d, numero):
    diapo = base_contenido(prs, d, numero)
    w_izq = ANCHO - W_DER_ETL - HUECO_ETL
    alto_sup = ETL['alto_sup']
    caja(diapo, X0 + 0.1, Y0 + 1, w_izq, alto_sup, parrafos_paso(d), 1.0, nombre='Problema y método')
    tarjeta_ejemplos(diapo, d, X1 - W_DER_ETL, Y0, W_DER_ETL, alto_sup)
    alto_cod, _ = caja_codigo(diapo, d, X0, ETL['y_codigo'], ANCHO, TAM_CODIGO_ETL)
    y = ETL['y_codigo'] + alto_cod + SEPARACION_ETL
    pars_res = parrafos_destacado('RESULTADO', d['resultado'])
    alto_res = alto_bloque(pars_res, ANCHO, MARGEN_DESTACADO)
    if y + alto_res > Y1 + 0.5:
        PROBLEMAS.append(f"{d['id']}: el resultado necesita {alto_res:.0f} pt y hay {Y1 - y:.0f}")
    bloque(diapo, X0, y, ANCHO, Y1 - y, VERDE, pars_res, 1.0, MARGEN_DESTACADO, ancla='m', nombre='Resultado')
    if d.get('figura'):
        PROBLEMAS.append(f"{d['id']}: la figura opcional de paso_etl no se dibuja en este diseño")
    return diapo


W_FIG_TALLER = 327.0


def nota_sintetica(diapo):
    caja(diapo, X1 - 320, y_texto(73.0, 10) + 1.0, 320, 14,
         [P(R(NOTA_SINTETICOS, 8.6, SECUNDARIO, cursiva=True), alin='r')], nombre='Nota de datos sintéticos')


def grafica_taller(prs, d, numero):
    diapo = base_contenido(prs, d, numero)
    nota_sintetica(diapo)
    alto_cod, _ = caja_codigo(diapo, d, X0, Y0, ANCHO)
    y = Y0 + alto_cod + 16
    alto_disp = Y1 - y
    with Image.open(FUENTES / d['figura']) as im:
        proporcion = im.width / im.height
    w_fig = W_FIG_TALLER
    h_fig = w_fig / proporcion
    if h_fig > alto_disp + 0.5:
        AVISOS.append(f"{d['id']}: la figura se achicó a {alto_disp * proporcion:.0f} pt de ancho para caber")
        h_fig = alto_disp
        w_fig = h_fig * proporcion
    fx, fy, fw, fh = figura(diapo, d['figura'], X1 - w_fig, Y1 - h_fig, w_fig, h_fig, 'r', 't', d['pregunta'])
    w_izq = X1 - W_FIG_TALLER - 18 - X0
    pars_dest = parrafos_destacado('PREGUNTA PARA EL GRUPO', d['pregunta_grupo'], 12.0)
    alto_dest = alto_bloque(pars_dest, w_izq, (14, 11, 14, 11))
    izq = [P(R(d['pregunta'], 14, VERDE), despues=7, inter=1.0),
           P(*mixto(d['lectura'], 11), inter=1.06)]
    alto_izq_disp = Y1 - y - alto_dest - 12
    escala = ajustar(izq, w_izq, alto_izq_disp, d['id'])
    caja(diapo, X0 + 0.1, y, w_izq, alto_izq_disp, izq, escala, nombre='Pregunta y lectura')
    bloque(diapo, X0, Y1 - alto_dest, w_izq, alto_dest, VERDE, pars_dest, 1.0, (14, 11, 14, 11),
           nombre='Pregunta para el grupo')
    return diapo


def lista_numerada(prs, d, numero):
    diapo = base_contenido(prs, d, numero)
    items = d['items']
    n = len(items)
    paso = min(46.0, (Y1 - Y0) / n)
    x_num = X0 + 16
    x_txt = X0 + 40
    w_txt = X1 - x_txt - 16
    for k, it in enumerate(items):
        y = Y0 + k * paso
        if k % 2 == 0:
            rectangulo(diapo, X0, y, ANCHO, paso - 2, TARJETA, REDONDEO, nombre=f'Banda {k + 1}')
        pars = [P(R(it['negrita'] + ' ', 12, TEXTO, True), *mixto(it['texto'], 12))]
        escala = ajustar(pars, w_txt, paso - 8, f"{d['id']} ítem {k + 1}")
        caja(diapo, x_num, y, 24, paso - 2, [P(R(f'{k + 1}.', 12, VERDE, True))], escala, 'm', nombre=f'Número {k + 1}')
        caja(diapo, x_txt, y, w_txt, paso - 2, pars, escala, 'm', nombre=f'Ítem {k + 1}')
    return diapo


def cierre(prs, d, numero):
    diapo = nueva(prs, 'fondo_cierre.png')
    caja(diapo, X0 + 0.1, y_texto(107.8, 30), ANCHO, 40, [P(R(d['titulo'], 30, CREMA))], nombre='Título')
    caja(diapo, X0 + 0.1, y_texto(162.9, 9.6), ANCHO, 14, [etiqueta_runs(d['etiqueta'], 9.6, DORADO)],
         nombre='Etiqueta')
    pars = [P(R(t, 9.2, CREMA), despues=16.1 - 9.2 * ALTO_LINEA['Calibri']) for t in d['lineas']]
    alto = alto_texto(pars, ANCHO)
    if 191.0 + alto + 30 > 305:
        PROBLEMAS.append(f"{d['id']}: las líneas del cierre bajan demasiado")
    caja(diapo, X0 + 0.1, y_texto(191.0, 9.2), ANCHO, alto + 4, pars, nombre='Líneas')
    y_firma = 191.0 + alto + 24
    caja(diapo, X0 + 0.1, y_texto(y_firma, 9.6), ANCHO, 14, [P(R(d['firma'], 9.6, CREMA, cursiva=True))],
         nombre='Firma')
    return diapo


LAYOUTS = {
    'portada': portada,
    'seccion': seccion,
    'tarjetas': tarjetas,
    'cifras': cifras,
    'texto_figura': texto_figura,
    'tabla': tabla,
    'flujo': flujo,
    'paso_etl': paso_etl,
    'grafica_taller': grafica_taller,
    'lista_numerada': lista_numerada,
    'cierre': cierre,
}


def propiedades(prs, total):
    p = prs.core_properties
    p.title = TITULO_DOC
    p.author = AUTOR
    p.last_modified_by = AUTOR
    p.subject = 'Minicurso de 1 h 30 min: limpieza (ETL) y gráficas con pandas en Google Colab'
    p.keywords = 'análisis de datos; pandas; ETL; limpieza de datos; gráficas; Python; UNAL'
    p.category = 'Minicurso'
    p.comments = (f'{total} diapositivas. Universidad Nacional de Colombia, Facultad de Minas, Sede Medellín. '
                  'Los datos de «Pan de Datos» son sintéticos, creados para el curso.')
    p.language = 'es-CO'
    p.revision = 1
    ahora = datetime.datetime(2026, 10, 18, 12, 0, 0)
    p.created = ahora
    p.modified = ahora


def main():
    diapositivas = []
    enlace = '`' + (ENLACE_TALLER.strip() or MARCADOR_ENLACE) + '`'
    con_enlace = {}
    for archivo in CONTENIDO:
        for d in json.loads(archivo.read_text(encoding='utf-8')):
            texto = json.dumps(d, ensure_ascii=False)
            if '{enlace}' in texto:
                d = json.loads(texto.replace('{enlace}', enlace))
                con_enlace[d['id']] = len(diapositivas) + 1
            diapositivas.append(d)
    numeros = ', '.join(str(n) for n in con_enlace.values())
    if not ENLACE_TALLER.strip():
        AVISOS.append(f'falta el enlace real del taller: las diapositivas {numeros} muestran «{MARCADOR_ENLACE}» '
                      '(llenar ENLACE_TALLER en construir_presentacion.py)')
    preparar_pasos(diapositivas)
    ids = [d['id'] for d in diapositivas]
    if len(ids) != len(set(ids)):
        PROBLEMAS.append('hay ids repetidos')
    prs = Presentation()
    prs.slide_width = Pt(720)
    prs.slide_height = Pt(405)
    for numero, d in enumerate(diapositivas, start=1):
        if d['layout'] not in LAYOUTS:
            PROBLEMAS.append(f"{d['id']}: layout desconocido {d['layout']}")
            continue
        diapo = LAYOUTS[d['layout']](prs, d, numero)
        if not d.get('notas'):
            PROBLEMAS.append(f"{d['id']}: no tiene notas del orador")
        notas(diapo, d)
    if ENLACE_TALLER.strip() and any(p.startswith(tuple(con_enlace)) for p in PROBLEMAS):
        PROBLEMAS.append(f'el enlace del taller ({len(ENLACE_TALLER.strip())} caracteres) no cabe en las diapositivas '
                         f'{numeros}: usa un enlace corto, de un acortador como bit.ly o tinyurl.com')
    propiedades(prs, len(prs.slides))
    SALIDA.parent.mkdir(exist_ok=True)
    prs.save(SALIDA)
    print(f'{len(prs.slides)} diapositivas escritas en {SALIDA}')
    if AVISOS:
        print('AVISOS:')
        print('\n'.join(f'- {a}' for a in AVISOS))
    if PROBLEMAS:
        print('PROBLEMAS:')
        print('\n'.join(f'- {p}' for p in PROBLEMAS))
        sys.exit(1)


if __name__ == '__main__':
    main()
